# ============================================================
#  TeamX Digital Library — Data Generation Script
#  Author: Roselyn G. Polanco González (2309184)
#  Role:   Data Generation & Instrumentation
#
#  Generates realistic sample data (CSV + JSON + Excel) that
#  supports all five team research questions:
#    [Arturo] search_events  — filter_used vs. clicked
#    [Diego]  book_views      — subject traffic by career
#    [Rous]   book_subjects   — single vs. multi-tag views
#    [Rodri]  sessions/saves  — browsing depth vs. save
#    [Mau]    sessions        — duration vs. distinct subjects
# ============================================================

import csv
import json
import os
import random
from datetime import datetime, timedelta

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_OK = True
except ImportError:
    EXCEL_OK = False
    print("[WARN] openpyxl missing — Excel export skipped. pip install openpyxl")

# ── Reproducibility ──────────────────────────────────────────
random.seed(42)

# ── Paths ────────────────────────────────────────────────────
BASE = os.path.dirname(os.path.abspath(__file__))
CSV_DIR  = os.path.abspath(os.path.join(BASE, "../data/csv"))
JSON_DIR = os.path.abspath(os.path.join(BASE, "../data/json"))
XLS_DIR  = os.path.abspath(os.path.join(BASE, "../data/excel"))
for d in (CSV_DIR, JSON_DIR, XLS_DIR):
    os.makedirs(d, exist_ok=True)

# ── Reference dimensions ─────────────────────────────────────
CAREERS = ["Data Engineering", "Robotics", "Biotechnology", "Embedded Systems"]
SUBJECTS = ["Mathematics", "Programming", "Databases", "Statistics",
            "Electronics", "Biology", "Machine Learning", "Ethics"]

# subjects each career tends to read (creates the signal for Diego's question)
CAREER_BIAS = {
    "Data Engineering": ["Programming", "Databases", "Statistics", "Machine Learning"],
    "Robotics":         ["Electronics", "Programming", "Mathematics"],
    "Biotechnology":    ["Biology", "Statistics", "Ethics"],
    "Embedded Systems": ["Electronics", "Programming", "Mathematics"],
}

N_STUDENTS = 200
N_BOOKS    = 120
N_VIEWS    = 5000
N_SEARCHES = 1500
N_SESSIONS = 800

BASE_DATE = datetime(2025, 5, 1)


def rand_dt():
    return BASE_DATE + timedelta(
        days=random.randint(0, 27),
        hours=random.randint(8, 21),
        minutes=random.randint(0, 59),
    )


# ── 1. Students ──────────────────────────────────────────────
students = []
for sid in range(1, N_STUDENTS + 1):
    students.append({
        "student_id": sid,
        "control_num": f"23{random.randint(10000, 99999)}",
        "full_name": f"Student {sid:03d}",
        "email": f"student{sid:03d}@uni.mx",
        "career": random.choice(CAREERS),
        "semester": random.randint(1, 8),
    })
career_of = {s["student_id"]: s["career"] for s in students}


# ── 2. Subjects ──────────────────────────────────────────────
subjects = [{"subject_id": i + 1, "name": name} for i, name in enumerate(SUBJECTS)]
subject_id_of = {s["name"]: s["subject_id"] for s in subjects}


# ── 3. Books + book_subjects (single vs multi tag → Rous) ────
books = []
book_subjects = []
for bid in range(1, N_BOOKS + 1):
    books.append({
        "book_id": bid,
        "title": f"Book Title {bid:03d}",
        "author": f"Author {random.randint(1, 60):02d}",
        "year": random.randint(2005, 2024),
        "description": "Sample academic book description.",
    })
    # ~45% single-tag, ~55% multi-tag (2-4 subjects)
    n_tags = 1 if random.random() < 0.45 else random.randint(2, 4)
    tags = random.sample(SUBJECTS, n_tags)
    for t in tags:
        book_subjects.append({"book_id": bid, "subject_id": subject_id_of[t]})

# map book → its subjects (names) for view generation
book_tag_names = {}
for bs in book_subjects:
    book_tag_names.setdefault(bs["book_id"], []).append(
        next(s["name"] for s in subjects if s["subject_id"] == bs["subject_id"])
    )
# multi-tagged books get a popularity boost (the effect Rous is testing)
book_weight = {}
for b in books:
    bid = b["book_id"]
    n = len(book_tag_names.get(bid, []))
    book_weight[bid] = 1.0 + 0.6 * (n - 1)   # more tags → more likely to be viewed


# ── 4. Book views (career traffic → Diego; tag effect → Rous) ─
book_ids = [b["book_id"] for b in books]
weights = [book_weight[b] for b in book_ids]
views = []
for vid in range(1, N_VIEWS + 1):
    stu = random.choice(students)
    career = stu["career"]
    # pick a book; 65% of time steer toward a book tagged with a career-biased subject
    if random.random() < 0.65:
        liked = CAREER_BIAS[career]
        candidates = [bid for bid in book_ids
                      if set(book_tag_names.get(bid, [])) & set(liked)]
        book_id = random.choice(candidates) if candidates else random.choices(book_ids, weights)[0]
    else:
        book_id = random.choices(book_ids, weights=weights, k=1)[0]
    views.append({
        "view_id": vid,
        "book_id": book_id,
        "student_id": stu["student_id"],
        "career": career,
        "viewed_at": rand_dt().strftime("%Y-%m-%d %H:%M:%S"),
    })


# ── 5. Search events (filter vs click → Arturo) ──────────────
search_events = []
for eid in range(1, N_SEARCHES + 1):
    stu = random.choice(students)
    filter_used = random.random() < 0.5
    # filtered searches convert better (the hypothesis Arturo will test)
    p_click = 0.62 if filter_used else 0.41
    clicked = random.random() < p_click
    search_events.append({
        "event_id": eid,
        "student_id": stu["student_id"],
        "query": random.choice(SUBJECTS).lower(),
        "filter_used": filter_used,
        "clicked": clicked,
        "searched_at": rand_dt().strftime("%Y-%m-%d %H:%M:%S"),
    })


# ── 6. Sessions (duration vs subjects → Mau; depth vs save → Rodri)
sessions = []
session_subjects = []
saved_books = []
save_id = 1
for sess_id in range(1, N_SESSIONS + 1):
    stu = random.choice(students)
    start = rand_dt()
    n_subjects = random.randint(1, 6)
    # duration correlates with number of subjects visited (Mau's hypothesis)
    duration = round(n_subjects * random.uniform(4, 9) + random.uniform(-3, 3), 2)
    duration = max(2.0, duration)
    end = start + timedelta(minutes=duration)
    books_browsed = random.randint(1, 25)
    # deeper browsing → higher save probability (Rodri's hypothesis)
    p_save = min(0.92, 0.15 + 0.03 * books_browsed)
    saved_final = random.random() < p_save

    sessions.append({
        "session_id": sess_id,
        "student_id": stu["student_id"],
        "started_at": start.strftime("%Y-%m-%d %H:%M:%S"),
        "ended_at": end.strftime("%Y-%m-%d %H:%M:%S"),
        "duration_minutes": duration,
        "books_browsed": books_browsed,
        "saved_final": saved_final,
    })
    for sub in random.sample(SUBJECTS, n_subjects):
        session_subjects.append({
            "session_id": sess_id,
            "subject_id": subject_id_of[sub],
        })
    if saved_final:
        saved_books.append({
            "save_id": save_id,
            "student_id": stu["student_id"],
            "book_id": random.choice(book_ids),
            "browsing_depth_at_save": books_browsed,
            "saved_at": end.strftime("%Y-%m-%d %H:%M:%S"),
        })
        save_id += 1


# ── Writers ──────────────────────────────────────────────────
def write_csv(name, rows, fields):
    path = os.path.join(CSV_DIR, name)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"[CSV]  {name:<22} {len(rows):>5} rows")


def write_json(name, rows):
    path = os.path.join(JSON_DIR, name)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    print(f"[JSON] {name:<22} {len(rows):>5} records")


def write_excel():
    if not EXCEL_OK:
        return
    wb = openpyxl.Workbook()

    # Sheet 1 — book tag summary (Rous)
    ws1 = wb.active
    ws1.title = "Book Tag Summary"
    view_count = {}
    for v in views:
        view_count[v["book_id"]] = view_count.get(v["book_id"], 0) + 1
    headers = ["book_id", "title", "num_tags", "tags", "total_views"]
    _style_header(ws1, headers)
    for b in books:
        bid = b["book_id"]
        tags = book_tag_names.get(bid, [])
        ws1.append([bid, b["title"], len(tags), ", ".join(tags), view_count.get(bid, 0)])

    # Sheet 2 — session stats (Mau / Rodri)
    ws2 = wb.create_sheet("Session Stats")
    subj_per_session = {}
    for ss in session_subjects:
        subj_per_session[ss["session_id"]] = subj_per_session.get(ss["session_id"], 0) + 1
    headers2 = ["session_id", "career", "duration_minutes", "distinct_subjects",
                "books_browsed", "saved_final"]
    _style_header(ws2, headers2)
    for s in sessions:
        ws2.append([
            s["session_id"], career_of[s["student_id"]], s["duration_minutes"],
            subj_per_session.get(s["session_id"], 0), s["books_browsed"], s["saved_final"],
        ])

    # Sheet 3 — search conversion (Arturo)
    ws3 = wb.create_sheet("Search Conversion")
    _style_header(ws3, ["group", "searches", "clicks", "conversion_rate"])
    for grp, flag in [("with_filter", True), ("no_filter", False)]:
        subset = [e for e in search_events if e["filter_used"] == flag]
        clicks = sum(1 for e in subset if e["clicked"])
        rate = round(clicks / len(subset), 3) if subset else 0
        ws3.append([grp, len(subset), clicks, rate])

    for ws in wb.worksheets:
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col) + 2
            ws.column_dimensions[col[0].column_letter].width = min(width, 40)

    path = os.path.join(XLS_DIR, "library_analytics.xlsx")
    wb.save(path)
    print(f"[XLSX] library_analytics.xlsx   (3 sheets)")


def _style_header(ws, headers):
    fill = PatternFill("solid", fgColor="2563EB")
    font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


# ── Main ─────────────────────────────────────────────────────
if __name__ == "__main__":
    print("=== TeamX Library Data Generator (Rous — 2309184) ===\n")

    write_csv("students.csv", students,
              ["student_id", "control_num", "full_name", "email", "career", "semester"])
    write_csv("books.csv", books,
              ["book_id", "title", "author", "year", "description"])
    write_csv("subjects.csv", subjects, ["subject_id", "name"])
    write_csv("book_subjects.csv", book_subjects, ["book_id", "subject_id"])
    write_csv("book_views.csv", views,
              ["view_id", "book_id", "student_id", "career", "viewed_at"])
    write_csv("search_events.csv", search_events,
              ["event_id", "student_id", "query", "filter_used", "clicked", "searched_at"])
    write_csv("sessions.csv", sessions,
              ["session_id", "student_id", "started_at", "ended_at",
               "duration_minutes", "books_browsed", "saved_final"])
    write_csv("session_subjects.csv", session_subjects, ["session_id", "subject_id"])
    write_csv("saved_books.csv", saved_books,
              ["save_id", "student_id", "book_id", "browsing_depth_at_save", "saved_at"])

    print()
    write_json("students.json", students)
    write_json("books.json", books)
    write_json("subjects.json", subjects)
    write_json("book_views.json", views)

    print()
    write_excel()

    print("\n✅ All sample data generated.")
